"""The Dashboard tab's API: saved studies, loading a study from the fake Curious, and the four views."""

from __future__ import annotations

import csv
import io
import json
import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from curious_download import cli
from curious_download.client import CuriousClient
from curious_download.dashboard import loader
from curious_download.webapp.server import create_app

from . import fake_dashboard as fake
from .fake_dashboard import (
    BASELINE,
    CONSENT,
    DEMOGRAPHICS,
    EMA,
    EVENING,
    INTAKE,
    MORNING,
    OWNER_ID,
    SLEEP,
    WEEKLY_FLOW,
    day,
    first_weekday_from,
)

TOKEN = "test-token"
MORNING_KEY = f"{EMA}:activity:{MORNING}"
EVENING_KEY = f"{EMA}:activity:{EVENING}"
SLEEP_KEY = f"{EMA}:activity:{SLEEP}"
BASELINE_KEY = f"{EMA}:activity:{BASELINE}"
FLOW_KEY = f"{EMA}:flow:{WEEKLY_FLOW}"
CONSENT_KEY = f"{INTAKE}:activity:{CONSENT}"
DEMOGRAPHICS_KEY = f"{INTAKE}:activity:{DEMOGRAPHICS}"


@pytest.fixture
def make_client(monkeypatch, tmp_path):
    monkeypatch.setattr(cli, "SETTINGS_PATH", tmp_path / "settings.json")

    def build(server=None) -> TestClient:
        server = server or fake.FakeStudyServer()
        app = create_app(
            token=TOKEN,
            client_factory=lambda url: CuriousClient(url, transport=server.transport()),
            allowed_hosts={"testserver"},
        )
        client = TestClient(app)
        client.headers["X-Curious-Token"] = TOKEN
        client.fake = server  # type: ignore[attr-defined]
        response = client.post(
            "/api/login", json={"server": "https://api.curious.test", "email": fake.EMAIL, "password": fake.PASSWORD}
        )
        assert response.status_code == 200, response.text
        return client

    return build


def study_body(**overrides) -> dict:
    body = {
        "name": "EMA pilot",
        "view": "fixed",
        "applets": [
            {"ownerId": OWNER_ID, "appletId": EMA, "name": "EMA Study"},
            {"ownerId": OWNER_ID, "appletId": INTAKE, "name": "Intake"},
        ],
        "aliases": {"p002x": "P002"},
    }
    body.update(overrides)
    return body


def create(client: TestClient, **overrides) -> dict:
    response = client.post("/api/dashboard/studies", json=study_body(**overrides))
    assert response.status_code == 200, response.text
    return response.json()


def wait_for(client: TestClient, job_id: str, timeout: float = 30.0) -> dict:
    deadline = time.time() + timeout
    while time.time() < deadline:
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["status"] != "running":
            return job
        time.sleep(0.02)
    raise AssertionError("the load did not finish in time")


def load(client: TestClient, study_id: str) -> dict:
    started = client.post(f"/api/dashboard/studies/{study_id}/load")
    assert started.status_code == 200, started.text
    job = wait_for(client, started.json()["jobId"])
    assert job["status"] == "done", job["error"]
    return job


def view(client: TestClient, study_id: str, **params) -> dict:
    response = client.get(f"/api/dashboard/studies/{study_id}/view", params=params)
    assert response.status_code == 200, response.text
    return response.json()


def loaded_study(client: TestClient, **overrides) -> dict:
    study = create(client, **overrides)
    load(client, study["id"])
    return study


def cell(grid: dict, row_label: str, column_key: str) -> dict | None:
    row = next(r for r in grid["rows"] if r["label"] == row_label)
    index = next(i for i, c in enumerate(grid["columns"]) if c["key"] == column_key)
    return row["cells"][index]


def col(bucket, entity_key: str) -> str:
    return f"{bucket.isoformat() if hasattr(bucket, 'isoformat') else bucket}|{entity_key}"


# -- studies -------------------------------------------------------------------------------


def test_study_crud_and_validation(make_client):
    client = make_client()
    assert client.get("/api/dashboard/studies").json() == []

    missing = client.post("/api/dashboard/studies", json=study_body(applets=[]))
    assert missing.status_code == 400 and "at least one applet" in missing.json()["error"]
    bad_view = client.post("/api/dashboard/studies", json=study_body(view="pie"))
    assert bad_view.status_code == 400
    bad_date = client.post("/api/dashboard/studies", json=study_body(manual_starts={"P001": "yesterday"}))
    assert bad_date.status_code == 400 and "YYYY-MM-DD" in bad_date.json()["error"]

    study = create(client)
    assert study["sameServer"] and not study["loaded"] and study["server"] == "https://api.curious.test"
    assert study["created"] and study["id"]

    renamed = client.post("/api/dashboard/studies", json={**study, "name": "EMA pilot 2"}).json()
    assert renamed["id"] == study["id"] and renamed["name"] == "EMA pilot 2"
    assert [s["name"] for s in client.get("/api/dashboard/studies").json()] == ["EMA pilot 2"]

    assert client.get(f"/api/dashboard/studies/{study['id']}/view").status_code == 409, "no data loaded yet"
    assert client.delete(f"/api/dashboard/studies/{study['id']}").status_code == 200
    assert client.get("/api/dashboard/studies").json() == []
    assert client.get(f"/api/dashboard/studies/{study['id']}").status_code == 404
    assert client.get("/api/dashboard/studies/../etc").status_code in (400, 404)


def test_studies_need_a_login(make_client):
    client = make_client()
    client.post("/api/logout")
    assert client.get("/api/dashboard/studies").status_code == 401


# -- loading ----------------------------------------------------------------------------------


def test_load_reads_participants_schedules_and_completions(make_client):
    client = make_client()
    study = create(client)
    job = load(client, study["id"])
    # P001, P002 (p002x merged in), P003 (shell), P004 (removed, no data), P009 (removed, data kept)
    assert job["result"] == {
        "studyId": study["id"],
        "participants": 5,
        "completions": 16,
        "applets": 2,
        "activities": 7,
    }
    assert any("Intake" in warning and "schedule could not be read" in warning for warning in job["warnings"])

    status = client.get(f"/api/dashboard/studies/{study['id']}").json()
    assert status["loaded"] and status["loadingJobId"] is None
    data = status["data"]
    assert data["participants"] == 5 and data["completions"] == 16
    by_name = {applet["name"]: applet for applet in data["applets"]}
    assert by_name["EMA Study"]["scheduleAvailable"] and by_name["EMA Study"]["historyAvailable"]
    assert not by_name["Intake"]["scheduleAvailable"] and not by_name["Intake"]["historyAvailable"]
    # Reviewer-only, hidden, and flow-only activities get no column; Sleep does (it had its own schedule once).
    assert [e["label"] for e in data["entities"]] == [
        "Morning",
        "Evening",
        "Sleep",
        "Baseline survey",
        "Flow: Weekly flow",
        "Consent",
        "Demographics",
    ]
    assert "Load" not in status  # the whole config is echoed too
    assert status["aliases"] == {"p002x": "P002"}


def test_history_is_read_page_by_page(make_client, monkeypatch):
    monkeypatch.setattr(loader, "HISTORY_PAGE_SIZE", 3)
    client = make_client()
    study = loaded_study(client)
    history_calls = [r for r in client.fake.requests if r.url.path == f"/applets/{EMA}/events/history"]
    assert len(history_calls) == 3  # 8 rows, 3 per page
    grid = view(client, study["id"], view="fixed", first=day(-10).isoformat(), last=day(-7).isoformat())
    # The deleted "Sleep" event was in force from day -10 until it was deleted on day -8.
    assert cell(grid, "P001", col(day(-10), SLEEP_KEY))["s"] == "done"
    assert cell(grid, "P001", col(day(-9), SLEEP_KEY))["s"] == "missed"
    assert cell(grid, "P001", col(day(-7), SLEEP_KEY))["s"] == "none"


def test_data_from_limits_what_is_fetched(make_client):
    client = make_client()
    study = loaded_study(client, data_from=day(-8).isoformat())
    status = client.get(f"/api/dashboard/studies/{study['id']}").json()
    assert status["data"]["completions"] < 16
    exports = [r for r in client.fake.requests if r.url.path.endswith("/data")]
    assert all("fromDate" in str(r.url) for r in exports)


def test_saving_a_changed_applet_list_drops_the_loaded_data(make_client):
    client = make_client()
    study = loaded_study(client)
    same = client.post("/api/dashboard/studies", json={**study, "view": "cross"}).json()
    assert same["loaded"], "changing the dashboard type keeps the data"
    fewer = client.post("/api/dashboard/studies", json={**study, "applets": study["applets"][:1]}).json()
    assert not fewer["loaded"], "changing the applets needs a reload"


# -- fixed view ----------------------------------------------------------------------------------


def test_fixed_view_cells(make_client):
    client = make_client()
    study = loaded_study(client)
    grid = view(
        client, study["id"], view="fixed", first=day(-10).isoformat(), last=day(5).isoformat(), granularity="day"
    )

    assert grid["view"] == "fixed" and grid["params"]["first"] == day(-10).isoformat()
    assert [row["label"] for row in grid["rows"]] == ["P001", "P002", "P003"], "removed participants are left out"
    assert len(grid["groups"]) == 16 and len(grid["columns"]) == 16 * 7
    assert sum(1 for group in grid["groups"] if group["today"]) == 1

    # P001: in window, outside the window (counted as done, flagged), spread-to-next-day evening, legacy + repeat
    assert cell(grid, "P001", col(day(-10), MORNING_KEY)) == {
        "e": 1,
        "d": 1,
        "l": 0,
        "p": 0,
        "m": 0,
        "u": 0,
        "x": 0,
        "s": "done",
        "t": day(-10).isoformat(),
    }
    late = cell(grid, "P001", col(day(-9), MORNING_KEY))
    assert late["s"] == "done" and late["l"] == 1 and late["d"] == 1
    repeat = cell(grid, "P001", col(day(-8), MORNING_KEY))
    assert repeat["s"] == "done" and repeat["x"] == 1
    assert cell(grid, "P001", col(day(-10), EVENING_KEY))["s"] == "done"
    assert cell(grid, "P001", col(day(-9), EVENING_KEY))["s"] == "missed"
    assert cell(grid, "P001", col(day(-7), FLOW_KEY))["s"] == "done"
    assert cell(grid, "P001", col(day(-8), FLOW_KEY))["s"] == "none", "weekly: nothing expected on other days"
    assert cell(grid, "P001", col(day(3), MORNING_KEY))["s"] == "upcoming"
    # Baseline is always available: done once shows as "done without a schedule"
    assert cell(grid, "P001", col(day(-10), BASELINE_KEY))["s"] == "extra"

    # P002: an individual weekday-morning schedule replaces the default mornings only
    assert cell(grid, "P002", col(day(-8), MORNING_KEY))["s"] == "none", "individual schedule starts later"
    assert cell(grid, "P002", col(day(-8), EVENING_KEY))["s"] == "missed", "default evenings still apply"
    weekday = first_weekday_from(day(-6))
    assert cell(grid, "P002", col(weekday, MORNING_KEY))["s"] == "done"
    assert cell(grid, "P002", col(day(-7), FLOW_KEY))["s"] == "partial"
    assert cell(grid, "P002", col(day(-9), MORNING_KEY))["s"] == "none", "nothing expected before P002 was added"

    # P003 is a shell account in EMA only: nothing from before they were added, and no Intake cells
    assert cell(grid, "P003", col(day(-5), EVENING_KEY))["s"] == "none"
    assert cell(grid, "P003", col(day(-2), EVENING_KEY))["s"] == "done"
    assert cell(grid, "P003", col(day(-2), CONSENT_KEY)) is None

    p001 = next(row for row in grid["rows"] if row["label"] == "P001")
    assert p001["summary"]["pct"] is not None and p001["summary"]["last"] and p001["summary"]["next"]
    assert "Pat" in p001["sub"] and "EMA Study" in p001["sub"]
    assert grid["overall"]["participants"] == 3 and grid["totals"][0] is not None
    assert any("Intake" in note for note in grid["notes"])


def test_fixed_view_options(make_client):
    client = make_client()
    study = loaded_study(client)
    strict = view(
        client, study["id"], view="fixed", first=day(-9).isoformat(), last=day(-9).isoformat(), lateCountsAsDone="false"
    )
    late = cell(strict, "P001", col(day(-9), MORNING_KEY))
    assert late["s"] == "late" and late["d"] == 0 and late["l"] == 1

    weekly = view(
        client, study["id"], view="fixed", first=day(-10).isoformat(), last=day(5).isoformat(), granularity="week"
    )
    assert len(weekly["groups"]) in (3, 4) and all(group["label"].startswith("Week of") for group in weekly["groups"])

    only_morning = view(
        client, study["id"], view="fixed", first=day(-10).isoformat(), last=day(-9).isoformat(), entities=MORNING_KEY
    )
    assert only_morning["groups"] == [] and [c["entity"] for c in only_morning["columns"]] == [MORNING_KEY, MORNING_KEY]

    with_removed = view(
        client, study["id"], view="fixed", first=day(-6).isoformat(), last=day(-6).isoformat(), includeRemoved="true"
    )
    assert [row["label"] for row in with_removed["rows"]] == ["P001", "P002", "P003", "P004", "P009"]
    assert cell(with_removed, "P009", col(day(-6), MORNING_KEY))["s"] == "done"
    assert "removed" in next(row for row in with_removed["rows"] if row["label"] == "P009")["sub"]

    bad = client.get(f"/api/dashboard/studies/{study['id']}/view", params={"first": "soon"})
    assert bad.status_code == 400


# -- cross-sectional view ------------------------------------------------------------------------------


def test_cross_view(make_client):
    client = make_client()
    study = loaded_study(client, view="cross")
    grid = view(client, study["id"])
    assert grid["view"] == "cross" and [c["key"] for c in grid["columns"]][-2:] == [CONSENT_KEY, DEMOGRAPHICS_KEY]
    assert (
        cell(grid, "P001", CONSENT_KEY)["s"] == "done" and cell(grid, "P001", CONSENT_KEY)["t"] == day(-11).isoformat()
    )
    assert cell(grid, "P002", CONSENT_KEY)["s"] == "done", "p002x was merged into P002"
    assert cell(grid, "P002", DEMOGRAPHICS_KEY)["s"] == "none"
    assert cell(grid, "P003", CONSENT_KEY) is None
    # An always-available one-time activity is open for those who haven't done it
    assert cell(grid, "P001", BASELINE_KEY)["s"] == "done"
    assert cell(grid, "P002", BASELINE_KEY)["s"] == "due"
    assert cell(grid, "P002", FLOW_KEY)["s"] == "partial"
    p002 = next(row for row in grid["rows"] if row["label"] == "P002")
    assert p002["summary"]["e"] == 7 and p002["summary"]["d"] == 2
    assert grid["totals"][grid["columns"].index(next(c for c in grid["columns"] if c["key"] == CONSENT_KEY))]["d"] == 2
    assert {item["status"] for item in grid["legend"]} >= {"done", "partial", "none", "due"}


# -- rolling view ----------------------------------------------------------------------------------------


def test_rolling_view_lines_participants_up_by_their_start(make_client):
    client = make_client()
    study = loaded_study(client, view="rolling", study_days=5, start_rule="first_completion")
    grid = view(client, study["id"])
    assert grid["view"] == "rolling" and [g["label"] for g in grid["groups"]] == [f"Day {i}" for i in range(1, 6)]
    rows = {row["label"]: row for row in grid["rows"]}
    # P001's earliest completion of anything is the Intake consent, the day before EMA started.
    assert "first completion" in rows["P001"]["sub"] and day(-11).isoformat() in rows["P001"]["sub"]
    assert cell(grid, "P001", col(2, MORNING_KEY))["s"] == "done"
    assert cell(grid, "P001", col(3, MORNING_KEY))["l"] == 1
    # P002 (with p002x merged in) first completed the Intake consent on day -9; the flow on day -7 is day 3.
    assert day(-9).isoformat() in rows["P002"]["sub"]
    assert cell(grid, "P002", col(3, FLOW_KEY))["s"] == "partial"

    mornings = view(client, study["id"], startEntity=MORNING_KEY)
    rows = {row["label"]: row for row in mornings["rows"]}
    assert day(-10).isoformat() in rows["P001"]["sub"]
    assert cell(mornings, "P001", col(1, MORNING_KEY))["s"] == "done"
    assert first_weekday_from(day(-6)).isoformat() in rows["P002"]["sub"]

    joined = view(client, study["id"], startRule="joined")
    rows = {row["label"]: row for row in joined["rows"]}
    assert all("added to the applet" in row["sub"] for row in rows.values())
    assert day(-3).isoformat() in rows["P003"]["sub"]

    individual = view(client, study["id"], startRule="individual_schedule")
    rows = {row["label"]: row for row in individual["rows"]}
    assert "individual schedule" in rows["P002"]["sub"] and day(-6).isoformat() in rows["P002"]["sub"]
    assert "(fallback)" in rows["P001"]["sub"]

    weekly = view(client, study["id"], granularity="week", studyDays=10)
    assert [g["label"] for g in weekly["groups"]] == ["Week 1", "Week 2"]


def test_rolling_view_manual_starts(make_client):
    client = make_client()
    study = loaded_study(client, view="rolling", start_rule="manual", manual_starts={"P003": day(-1).isoformat()})
    grid = view(client, study["id"])
    rows = {row["label"]: row for row in grid["rows"]}
    assert "entered by hand" in rows["P003"]["sub"]
    assert "(fallback)" in rows["P001"]["sub"]


# -- individual view --------------------------------------------------------------------------------------


def test_individual_view_and_focus(make_client):
    client = make_client()
    study = loaded_study(client, view="individual")
    grid = view(client, study["id"], first=day(-10).isoformat(), last=day(-7).isoformat())
    assert grid["groups"] == [] and [c["key"] for c in grid["columns"]] == [day(d).isoformat() for d in range(-10, -6)]
    first_day = cell(grid, "P001", day(-10).isoformat())
    assert (first_day["e"], first_day["d"], first_day["x"], first_day["s"]) == (3, 3, 1, "done")
    rows = {row["label"]: row for row in grid["rows"]}
    assert rows["P002"]["sub"].startswith("individual schedule") and rows["P001"]["sub"].startswith("default schedule")

    focused = view(client, study["id"], first=day(-10).isoformat(), last=day(-7).isoformat(), focus="P001")
    assert focused["params"]["focusLabel"].startswith("P001")
    labels = [row["label"] for row in focused["rows"]]
    assert labels == [
        "EMA Study: Morning",
        "EMA Study: Evening",
        "EMA Study: Sleep",
        "EMA Study: Baseline survey",
        "EMA Study: Flow: Weekly flow",
        "Intake: Consent",
        "Intake: Demographics",
    ]
    assert cell(focused, "EMA Study: Flow: Weekly flow", day(-7).isoformat())["s"] == "done"
    assert client.get(f"/api/dashboard/studies/{study['id']}/view", params={"focus": "nobody"}).status_code == 400


# -- details and exports ------------------------------------------------------------------------------------


def test_cell_details(make_client):
    client = make_client()
    study = loaded_study(client)
    params = {"view": "fixed", "first": day(-10).isoformat(), "last": day(-7).isoformat()}
    response = client.get(
        f"/api/dashboard/studies/{study['id']}/cell", params={**params, "row": "P001", "col": col(day(-9), MORNING_KEY)}
    )
    assert response.status_code == 200, response.text
    (item,) = response.json()["items"]
    assert item["status"] == "late" and item["schedule"] == "default" and item["periodicity"] == "DAILY"
    assert item["windowStart"].endswith("08:00") and item["windowEnd"].endswith("12:00")
    (completion,) = item["completions"]
    assert completion["inWindow"] is False and completion["complete"] is True

    flow = client.get(
        f"/api/dashboard/studies/{study['id']}/cell", params={**params, "row": "P002", "col": col(day(-7), FLOW_KEY)}
    ).json()
    (item,) = flow["items"]
    (completion,) = item["completions"]
    assert completion["complete"] is False
    assert [(a["name"], a["done"]) for a in completion["activities"]] == [("Sleep", True), ("Mood", False)]

    empty = client.get(
        f"/api/dashboard/studies/{study['id']}/cell", params={**params, "row": "P003", "col": col(day(-9), MORNING_KEY)}
    ).json()
    assert empty["items"] == []


def test_csv_exports(make_client):
    client = make_client()
    study = loaded_study(client)
    params = {"view": "fixed", "first": day(-10).isoformat(), "last": day(-9).isoformat()}
    grid_csv = client.get(f"/api/dashboard/studies/{study['id']}/export", params={**params, "format": "grid"})
    assert grid_csv.status_code == 200 and grid_csv.headers["content-type"].startswith("text/csv")
    assert 'filename="dashboard_ema-pilot_fixed_grid_' in grid_csv.headers["content-disposition"]
    assert grid_csv.text.startswith("﻿participant,details,adherence,")
    lines = grid_csv.text.splitlines()
    assert len(lines) == 4 and lines[1].startswith("P001,") and "1/1 (done)" in lines[1]

    long_csv = client.get(f"/api/dashboard/studies/{study['id']}/export", params={**params, "format": "long"})
    rows = long_csv.text.lstrip("﻿").splitlines()
    assert (
        rows[0]
        == "participant,applet,activity,day,window_start,window_end,status,schedule,periodicity,completed_start,completed_end,in_window,complete,submit_id"
    )
    assert any(
        line.startswith(f"P001,EMA Study,Morning,{day(-9).isoformat()}")
        and ",late,default,DAILY," in line
        and ",no,yes," in line
        for line in rows
    )
    assert any(f"P002,EMA Study,Evening,{day(-9).isoformat()}" in line and ",missed," in line for line in rows)


# -- folder export --------------------------------------------------------------------------------------------


def read_rows(path: Path) -> list[dict]:
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf"), "CSVs carry a BOM"
    return list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))


def test_export_job_writes_json_and_csvs(make_client, tmp_path):
    client = make_client()
    study = loaded_study(client)
    params = {"view": "fixed", "first": day(-10).isoformat(), "last": day(-9).isoformat(), "granularity": "day"}
    started = client.post(
        f"/api/dashboard/studies/{study['id']}/export-job", json={"outputDir": str(tmp_path / "out"), "params": params}
    )
    assert started.status_code == 200, started.text
    job = wait_for(client, started.json()["jobId"])
    assert job["status"] == "done", job["error"]
    assert job["kind"] == "dashboard_export" and job["title"] == "Dashboard export — EMA pilot"
    assert any("Intake" in warning for warning in job["warnings"])

    result = job["result"]
    folder = Path(result["folder"])
    assert folder.name.startswith("dashboard_EMA_pilot_fixed_") and result["view"] == "fixed"
    assert sorted(p.name for p in folder.iterdir()) == [
        "dashboard.xlsx",
        "dashboard_cells.csv",
        "dashboard_grid.csv",
        "dashboard_grid.json",
        "export_info.json",
        "occurrences.csv",
        "participants.csv",
    ]
    assert result["files"] == sorted(p.name for p in folder.iterdir())
    assert result["counts"]["rows"] == 3 and result["counts"]["columns"] == 14

    # The JSON copy is exactly what the page renders for the same parameters.
    grid_json = json.loads((folder / "dashboard_grid.json").read_text())
    assert grid_json == view(client, study["id"], **params)

    from openpyxl import load_workbook

    workbook = load_workbook(folder / "dashboard.xlsx")
    assert workbook.sheetnames == ["Info", "Grid", "Cells", "Participants", "Occurrences"]
    info_sheet = workbook["Info"]
    fields = {row[0].value: row[1].value for row in info_sheet.iter_rows(min_row=2) if row[0].value}
    assert fields["study.name"] == "EMA pilot" and fields["view.kind"] == "fixed"
    assert fields["done"] == "Completed", "the status legend is on the Info sheet"
    cells_sheet = workbook["Cells"]
    assert cells_sheet.max_row == 1 + 3 * 14 and cells_sheet["A1"].value == "participant"
    assert isinstance(cells_sheet["G2"].value, int), "counts are numbers, not text"
    grid_sheet = workbook["Grid"]
    assert grid_sheet["A2"].value == "Participant" and grid_sheet["A3"].value == "P001"
    header = [grid_sheet.cell(row=2, column=c).value for c in range(4, 4 + 14)]
    assert header[:2] == ["Morning", "Evening"], "column labels without the applet prefix"
    done_cell = grid_sheet.cell(row=3, column=4)  # P001, Morning, day -10
    assert done_cell.value == "1/1" and done_cell.fill.fgColor.rgb.endswith("D7F1D7")
    participants_sheet = workbook["Participants"]
    assert participants_sheet["A2"].value == "P001" and participants_sheet.freeze_panes == "A2"

    info = json.loads((folder / "export_info.json").read_text())
    assert info["study"]["name"] == "EMA pilot" and info["study"]["aliases"] == {"p002x": "P002"}
    assert "server" not in info["study"] and info["server"] == "https://api.curious.test"
    assert info["view"]["kind"] == "fixed" and info["view"]["first"] == day(-10).isoformat()
    assert info["view"]["lateCountsAsDone"] is True and info["view"]["today"]
    assert info["data"]["participants"] == 5 and [a["name"] for a in info["data"]["applets"]] == ["EMA Study", "Intake"]
    assert info["summary"]["participants"] == 3 and info["summary"]["adherence"] is not None
    assert {s["status"] for s in info["statuses"]} >= {"done", "missed", "upcoming"}
    assert info["files"]["cells"] == "dashboard_cells.csv"
    assert set(info["columns"]) == set(info["files"].values()) - {"export_info.json"}
    assert info["files"]["workbook"] == "dashboard.xlsx" and info["files"]["gridJson"] == "dashboard_grid.json"

    cells = read_rows(folder / "dashboard_cells.csv")
    assert len(cells) == 3 * 14
    morning = next(r for r in cells if r["participant"] == "P001" and r["column"] == col(day(-10), MORNING_KEY))
    assert (morning["status"], morning["expected"], morning["done"], morning["activity"], morning["applet"]) == (
        "done",
        "1",
        "1",
        "Morning",
        "EMA Study",
    )
    assert (
        morning["column_group"].endswith(day(-10).strftime("%-d %b"))
        and morning["first_completion"] == day(-10).isoformat()
    )
    consent = next(r for r in cells if r["participant"] == "P003" and r["column"] == col(day(-10), CONSENT_KEY))
    assert consent["status"] == "not_in_applet" and consent["expected"] == ""

    participants = {r["participant"]: r for r in read_rows(folder / "participants.csv")}
    assert set(participants) == {"P001", "P002", "P003"}
    assert participants["P001"]["nickname"] == "Pat" and participants["P001"]["applets"] == "EMA Study, Intake"
    assert participants["P001"]["schedule"] == "default" and participants["P002"]["schedule"] == "individual"
    assert participants["P001"]["joined"] == day(-11).isoformat(), "added to Intake the day before EMA"
    assert participants["P001"]["removed"] == "no"
    assert participants["P001"]["expected"] and participants["P001"]["adherence"]
    assert participants["P001"]["start_date"] == "", "start dates belong to the rolling view"

    grid_lines = (folder / "dashboard_grid.csv").read_text(encoding="utf-8-sig").splitlines()
    assert grid_lines[0].startswith("participant,details,adherence,")
    occurrences = read_rows(folder / "occurrences.csv")
    assert any(r["participant"] == "P001" and r["activity"] == "Morning" and r["status"] == "late" for r in occurrences)

    jobs = client.get("/api/jobs").json()
    assert {job["kind"] for job in jobs} == {"dashboard", "dashboard_export"}


def test_export_job_rolling_and_focus_variants(make_client, tmp_path):
    client = make_client()
    study = loaded_study(client, view="rolling", study_days=3)
    started = client.post(
        f"/api/dashboard/studies/{study['id']}/export-job",
        json={"outputDir": str(tmp_path / "out"), "params": {"view": "rolling", "startRule": "joined"}},
    )
    job = wait_for(client, started.json()["jobId"])
    assert job["status"] == "done", job["error"]
    folder = Path(job["result"]["folder"])
    participants = {r["participant"]: r for r in read_rows(folder / "participants.csv")}
    assert participants["P003"]["start_date"] == day(-3).isoformat()
    assert participants["P003"]["start_source"] == "added to the applet"
    assert json.loads((folder / "export_info.json").read_text())["view"]["startRule"] == "joined"

    focused = client.post(
        f"/api/dashboard/studies/{study['id']}/export-job",
        json={
            "outputDir": str(tmp_path / "out"),
            "params": {
                "view": "individual",
                "focus": "P001",
                "first": day(-10).isoformat(),
                "last": day(-9).isoformat(),
            },
        },
    )
    job = wait_for(client, focused.json()["jobId"])
    assert job["status"] == "done", job["error"]
    folder = Path(job["result"]["folder"])
    assert (folder / "activities.csv").exists() and not (folder / "participants.csv").exists()
    activities = read_rows(folder / "activities.csv")
    assert {r["activity"] for r in activities} >= {"EMA Study: Morning", "Intake: Consent"}
    assert next(r for r in activities if r["activity"] == "EMA Study: Morning")["applet"] == "EMA Study"
    cells = read_rows(folder / "dashboard_cells.csv")
    assert "row" in cells[0] and "participant" not in cells[0]

    bad = client.post(
        f"/api/dashboard/studies/{study['id']}/export-job",
        json={"outputDir": str(tmp_path / "out"), "params": {"view": "fixed", "first": "soon"}},
    )
    assert bad.status_code == 400
    missing = client.post(f"/api/dashboard/studies/{study['id']}/export-job", json={"outputDir": "", "params": {}})
    assert missing.status_code == 400 and "folder" in missing.json()["error"]


def test_export_job_without_the_optional_files(make_client, tmp_path, monkeypatch):
    client = make_client()
    study = loaded_study(client)
    plain = client.post(
        f"/api/dashboard/studies/{study['id']}/export-job",
        json={"outputDir": str(tmp_path / "out"), "params": {"view": "cross"}, "workbook": False, "gridJson": False},
    )
    job = wait_for(client, plain.json()["jobId"])
    assert job["status"] == "done", job["error"]
    folder = Path(job["result"]["folder"])
    assert sorted(p.name for p in folder.iterdir()) == [
        "dashboard_cells.csv",
        "dashboard_grid.csv",
        "export_info.json",
        "occurrences.csv",
        "participants.csv",
    ]
    info = json.loads((folder / "export_info.json").read_text())
    assert "workbook" not in info["files"] and "gridJson" not in info["files"]

    # Without openpyxl the workbook is skipped with a warning, and the rest is still written.
    from curious_download.webapp import runners

    def no_openpyxl(*args, **kwargs):
        raise ImportError("No module named 'openpyxl'")

    monkeypatch.setattr(runners, "write_workbook", no_openpyxl)
    started = client.post(
        f"/api/dashboard/studies/{study['id']}/export-job",
        json={"outputDir": str(tmp_path / "out2"), "params": {"view": "cross"}},
    )
    job = wait_for(client, started.json()["jobId"])
    assert job["status"] == "done", job["error"]
    assert any("openpyxl" in warning for warning in job["warnings"])
    folder = Path(job["result"]["folder"])
    assert not (folder / "dashboard.xlsx").exists() and (folder / "dashboard_grid.json").exists()
    assert "workbook" not in json.loads((folder / "export_info.json").read_text())["files"]


def test_rows_carry_structured_metadata(make_client):
    client = make_client()
    study = loaded_study(client, view="rolling")
    grid = view(client, study["id"], view="rolling", startRule="joined")
    p002 = next(row for row in grid["rows"] if row["label"] == "P002")
    # P002 was added to Intake on day -9, a day before EMA, so that is their earliest "joined" date.
    assert p002["meta"]["schedule"] == "individual" and p002["meta"]["start"] == day(-9).isoformat()
    assert p002["meta"]["startSource"] == "added to the applet" and p002["meta"]["applets"] == ["EMA Study", "Intake"]
    cross = view(client, study["id"], view="cross")
    assert next(row for row in cross["rows"] if row["label"] == "P001")["meta"]["nickname"] == "Pat"


# -- demo server -------------------------------------------------------------------------------------------


def test_demo_server_says_it_is_fake(monkeypatch, tmp_path):
    from . import demo_server

    monkeypatch.setattr(cli, "SETTINGS_PATH", tmp_path / "settings.json")
    server = demo_server.DemoServer()
    app = create_app(
        token=TOKEN,
        client_factory=lambda url: CuriousClient(url, transport=server.transport()),
        allowed_hosts={"testserver"},
        banner=demo_server.BANNER,
    )
    client = TestClient(app)
    client.headers["X-Curious-Token"] = TOKEN
    assert "fake data" in client.get("/api/state").json()["banner"]
    real = client.post(
        "/api/login", json={"server": "https://api.real.test", "email": "someone@real.org", "password": "secret"}
    )
    assert real.status_code == 401
    assert "demo with fake data" in real.json()["error"] and fake.EMAIL in real.json()["error"]
    assert "someone@real.org" in real.json()["error"]
    ok = client.post("/api/login", json={"server": "https://x", "email": fake.EMAIL, "password": fake.PASSWORD})
    assert ok.status_code == 200 and ok.json()["ok"]
