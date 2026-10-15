"""The dashboard export as one Excel workbook: a sheet per file, plus the grid with its colours."""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path

from .views import LEGEND, cell_text

# Fill and ink per status, the same tints the page uses.
STATUS_COLOURS = {
    "done": ("D7F1D7", "0B5D0B"),
    "late": ("FBE3D8", "8A3A16"),
    "partial": ("FDEFC4", "6B4D00"),
    "missed": ("F7D4D4", "8C1F1F"),
    "due": ("DBE8FB", "1C4F99"),
    "upcoming": ("FFFFFF", "9AA0AD"),
    "extra": ("FFFFFF", "0B5D0B"),
    "none": ("F1F2F5", "9AA0AD"),
    "not_in_applet": ("FAFAFB", "9AA0AD"),
}

# Columns whose values are numbers; everything else (IDs, dates) stays text so "0012" is not turned into 12.
NUMERIC_COLUMNS = {"expected", "done", "late", "partly_done", "missed", "open_or_upcoming", "unscheduled", "adherence"}
_NUMBER = re.compile(r"-?\d+(\.\d+)?")


def _coerce(header: str, value: str):
    if header in NUMERIC_COLUMNS and value and _NUMBER.fullmatch(value):
        return float(value) if "." in value else int(value)
    return value


def _flatten(value, prefix: str = "") -> list[tuple[str, str]]:
    """(dotted name, text) pairs for a nested dict, for the Info sheet."""
    if isinstance(value, dict):
        rows: list[tuple[str, str]] = []
        for key, inner in value.items():
            rows.extend(_flatten(inner, f"{prefix}.{key}" if prefix else str(key)))
        return rows
    if isinstance(value, list):
        if all(not isinstance(item, (dict, list)) for item in value):
            return [(prefix, ", ".join(str(item) for item in value))]
        rows = []
        for index, item in enumerate(value, start=1):
            rows.extend(_flatten(item, f"{prefix}.{index}"))
        return rows
    return [(prefix, "" if value is None else str(value))]


def write_workbook(path: Path, *, info: dict, grid: dict, tables: dict[str, str]) -> None:
    """Write `path`: an Info sheet, a coloured Grid sheet, then one sheet per CSV in `tables`.

    Raises ImportError when openpyxl is not installed.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    bold = Font(bold=True)
    muted = Font(color="677084")

    # -- Info ---------------------------------------------------------------------
    sheet = workbook.active
    sheet.title = "Info"
    sheet.append(["Field", "Value"])
    sheet["A1"].font = bold
    sheet["B1"].font = bold
    skip = {"statuses", "files", "columns", "notes"}
    for name, text in _flatten({key: value for key, value in info.items() if key not in skip}):
        sheet.append([name, text])
    sheet.append([])
    sheet.append(["Status legend", ""])
    sheet.cell(row=sheet.max_row, column=1).font = bold
    for item in info.get("statuses") or []:
        sheet.append([item["status"], item["label"]])
        fill, ink = STATUS_COLOURS.get(item["status"], ("FFFFFF", "000000"))
        sheet.cell(row=sheet.max_row, column=1).fill = PatternFill("solid", fgColor=fill)
        sheet.cell(row=sheet.max_row, column=1).font = Font(color=ink, bold=True)
    sheet.append([])
    sheet.append(["Files", ""])
    sheet.cell(row=sheet.max_row, column=1).font = bold
    for name, description in (info.get("columns") or {}).items():
        sheet.append([name, description])
    if info.get("notes"):
        sheet.append([])
        sheet.append(["Notes", ""])
        sheet.cell(row=sheet.max_row, column=1).font = bold
        for note in info["notes"]:
            sheet.append(["", note])
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 100

    # -- Grid ----------------------------------------------------------------------
    sheet = workbook.create_sheet("Grid")
    view = grid["view"]
    focused = bool(grid["params"].get("focus"))
    head = ["Activity" if focused else "Participant", "Details", "Adherence"]
    groups = grid.get("groups") or []
    if groups:
        sheet.append(["", "", ""] + [""] * len(grid["columns"]))
        column = 4
        for group in groups:
            sheet.cell(row=1, column=column, value=group["label"]).font = bold
            sheet.cell(row=1, column=column).alignment = Alignment(horizontal="center")
            if group["span"] > 1:
                sheet.merge_cells(start_row=1, start_column=column, end_row=1, end_column=column + group["span"] - 1)
            column += group["span"]
    labels = []
    for column_info in grid["columns"]:
        label = column_info["label"]
        applet = column_info.get("applet") or ""
        if applet and label.startswith(f"{applet}: "):
            label = label[len(applet) + 2 :]
        labels.append(
            label
            if column_info.get("entity") is not None or not column_info.get("sub")
            else f"{label} {column_info['sub']}"
        )
    sheet.append(head + labels)
    header_row = sheet.max_row
    for index in range(1, len(head) + len(labels) + 1):
        cell = sheet.cell(row=header_row, column=index)
        cell.font = bold
        if index > len(head):
            cell.alignment = Alignment(textRotation=90, horizontal="center", vertical="bottom")
    sheet.row_dimensions[header_row].height = 110
    thin_green = Side(style="thin", color="0B5D0B")
    thin_gray = Side(style="thin", color="E2E5EC")

    def paint(cell, status: str) -> None:
        fill, ink = STATUS_COLOURS.get(status, ("FFFFFF", "1D2433"))
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(color=ink, bold=status not in ("none", "upcoming", "not_in_applet"))
        cell.alignment = Alignment(horizontal="center")
        if status == "extra":
            cell.border = Border(left=thin_green, right=thin_green, top=thin_green, bottom=thin_green)
        elif status == "upcoming":
            cell.border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    def write_row(label: str, sub: str, pct, cells: list[dict | None], *, strong: bool = False) -> None:
        sheet.append(
            [label, sub, "" if pct is None else pct] + [cell_text(cell, view) if cell else "" for cell in cells]
        )
        row = sheet.max_row
        sheet.cell(row=row, column=1).font = bold if strong else Font()
        sheet.cell(row=row, column=2).font = muted
        if pct is not None:
            sheet.cell(row=row, column=3).number_format = "0%"
        for offset, cell in enumerate(cells):
            paint(sheet.cell(row=row, column=len(head) + 1 + offset), cell["s"] if cell else "not_in_applet")

    for row in grid["rows"]:
        write_row(row["label"], row.get("sub") or "", row["summary"].get("pct"), row["cells"])
    write_row(
        "All activities" if focused else "All participants",
        "",
        (grid.get("overall") or {}).get("pct"),
        grid.get("totals") or [],
        strong=True,
    )
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=len(head) + 1)
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 36
    sheet.column_dimensions["C"].width = 10
    for index in range(len(head) + 1, len(head) + len(labels) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 7 if view != "cross" else 12
    sheet.append([])
    sheet.append(["Legend"])
    sheet.cell(row=sheet.max_row, column=1).font = bold
    legend = grid.get("legend") or LEGEND
    for item in legend:
        sheet.append([item["status"], item["label"]])
        paint(sheet.cell(row=sheet.max_row, column=1), item["status"])

    # -- one sheet per CSV -----------------------------------------------------------
    for title, text in tables.items():
        sheet = workbook.create_sheet(title[:31])
        rows = list(csv.reader(io.StringIO(text)))
        if not rows:
            continue
        header = rows[0]
        sheet.append(header)
        for cell in sheet[1]:
            cell.font = bold
        for values in rows[1:]:
            sheet.append([_coerce(name, value) for name, value in zip(header, values, strict=False)])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, name in enumerate(header, start=1):
            longest = max(
                (len(str(values[index - 1])) for values in rows[:200] if index - 1 < len(values)), default=len(name)
            )
            sheet.column_dimensions[get_column_letter(index)].width = min(max(len(name), longest) + 2, 60)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
